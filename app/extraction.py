from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import io


def extract_pdf_text(file) -> dict:
    """
    Extracts text from each page of a PDF using pypdf.
    Returns a dict like {1: "page 1 text", 2: "page 2 text", ...}
    """
    reader = PdfReader(file)
    return {
        i + 1: page.extract_text() or ""
        for i, page in enumerate(reader.pages)
    }


def extract_docx_text(file) -> dict:
    """
    Extract each paragraph as a separate section.
    Returns: {1: "Paragraph 1 text", 2: "Paragraph 2 text", ...}
    """
    doc = Document(file)
    paragraphs = [
        p.text.strip() for p in doc.paragraphs if p.text.strip()
    ]
    return {
        i + 1: para for i, para in enumerate(paragraphs)
    }


def extract_pptx_text(file) -> dict:
    prs = Presentation(file)
    slide_texts = {}
    for i, slide in enumerate(prs.slides):
        text = " ".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text")
        )
        slide_texts[i + 1] = text
    return slide_texts


def extract_text_by_type(file, filetype) -> dict:
    if filetype == "pdf":
        return extract_pdf_text(file)
    elif filetype == "docx":
        return extract_docx_text(file)
    elif filetype == "pptx":
        return extract_pptx_text(file)
    else:
        raise ValueError("Unsupported file type")